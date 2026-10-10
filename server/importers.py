"""Bounded, side-effect-free imports with exact source locations.

Dates are never manufactured: a year in a sheet name is a reporting period,
not an observation timestamp. Aggregate figures and documents stay context;
only a source's actual feedback column is treated as customer testimony.
"""
from __future__ import annotations

import csv
import io
import json
import math
import re
from datetime import date, datetime
from html.parser import HTMLParser
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile

MAX_FILE_BYTES = 32 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED = 128 * 1024 * 1024
MAX_ZIP_MEMBERS = 10_000
MAX_ROWS = 100_000
MAX_COLUMNS = 1_000
MAX_CELLS = 2_000_000
MAX_RECORDS = 100_000
MAX_TEXT_CHARS = 20_000_000
MAX_METADATA_CHARS = 32_000_000
MAX_CELL_CHARS = 200_000
MAX_PDF_PAGES = 1_000
MAX_PDF_STREAM_BYTES = 16 * 1024 * 1024

_VEHICLE = re.compile(r"(?<![A-Z0-9])([FG]\d{2})(?![A-Z0-9])", re.I)
_MARKETS = {"USA": "US", "UNITED STATES": "US", "UNITED STATES OF AMERICA": "US",
            "CHINA": "CN", "EUROPE": "EU", "UNITED KINGDOM": "GB", "UK": "GB",
            "GERMANY": "DE", "REST OF WORLD": "ROW"}
_FEEDBACK_KEYS = ("customerfeedback", "feedback", "comment", "review")
_TEXT_KEYS = _FEEDBACK_KEYS + ("content", "text", "description")
_VEHICLE_KEYS = ("vehicle", "derivateecode", "ecode", "modelcode", "vehiclecode")
_MARKET_KEYS = ("marketcode", "countrycode", "country", "market")
_DATE_KEYS = ("observedat", "date", "feedbackdate", "observationdate", "publishedat", "timestamp")


def normalize_market(value: str) -> str:
    value = value.strip().upper()
    return _MARKETS.get(value, value)


def _key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    result = str(value).strip()
    if len(result) > MAX_CELL_CHARS:
        raise ValueError("A source value exceeds the text size limit")
    return result


def _json_value(value: object, depth: int = 0) -> object:
    if depth > 12:
        raise ValueError("Source metadata is nested too deeply")
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Source contains a non-finite number")
    if isinstance(value, str):
        _text(value)
    if isinstance(value, dict):
        result = {}
        for key, item in value.items():
            _text(key)
            result[str(key)] = _json_value(item, depth + 1)
        return result
    if isinstance(value, (list, tuple)):
        return [_json_value(v, depth + 1) for v in value]
    return value


def _lookup(columns: dict, keys: tuple[str, ...], *, text_only: bool = False) -> object:
    for key in keys:
        for name, value in columns.items():
            if text_only and not isinstance(value, str):
                continue
            if _key(name) == key and value is not None and (not isinstance(value, str) or value.strip()):
                return value
    return None


class _Records(list):
    """Check output budgets before accumulating each extracted record."""

    def __init__(self):
        super().__init__()
        self.text_chars = self.metadata_chars = 0

    def append(self, record: dict):
        self.text_chars += len(record["content"])
        self.metadata_chars += len(json.dumps(record["metadata"], ensure_ascii=False, allow_nan=False))
        if len(self) >= MAX_RECORDS or self.text_chars > MAX_TEXT_CHARS or self.metadata_chars > MAX_METADATA_CHARS:
            raise ValueError("File exceeds record, extracted text or metadata limits")
        super().append(record)

    def extend(self, records):
        for record in records:
            self.append(record)


def _date_value(value: object) -> str | None:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if not isinstance(value, str) or not value.strip():
        return None
    text = value.strip()
    try:
        # Accept explicit ISO dates/timestamps; keep partial or ambiguous dates raw.
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
            return date.fromisoformat(text).isoformat()
        if re.match(r"\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}", text):
            return datetime.fromisoformat(text.replace("Z", "+00:00")).isoformat()
    except ValueError:
        pass
    return None


def _vehicles(text: object) -> list[str]:
    return list(dict.fromkeys(code.upper() for code in _VEHICLE.findall(_text(text))))


def _sheet_market(sheet: str) -> str | None:
    tokens = re.split(r"[_\s-]+", sheet.upper())
    markets = {normalize_market(token) for token in tokens if token in {"US", "USA", "EU", "CN", "DE", "GB", "UK", "ROW"}}
    return markets.pop() if len(markets) == 1 else None


def _base_record(content: str, location: dict, metadata: dict, vehicle: str, market: str,
                 vehicles: list[str] | None = None, explicit_market: str | None = None,
                 observed_at: str | None = None) -> dict:
    metadata["scope_basis"] = {"vehicle": "source" if vehicles else "requested", "market": "source" if explicit_market else "requested"}
    if vehicles:
        metadata["vehicle_scope"] = vehicles
    if explicit_market:
        metadata["source_market"] = normalize_market(explicit_market)
    return {"content": content, "location": location, "metadata": metadata,
            "vehicle": vehicle if not vehicles or vehicle in vehicles else vehicles[0],
            "market": normalize_market(explicit_market) if explicit_market else market,
            "observed_at": observed_at}


def _columns(headers: list, values: tuple | list) -> dict:
    result = {}
    for i in range(max(len(headers), len(values))):
        value = values[i] if i < len(values) else None
        name = _text(headers[i]) if i < len(headers) else ""
        name = name or f"column_{i + 1}"
        unique = name
        suffix = 2
        while unique in result:
            unique = f"{name} ({suffix})"
            suffix += 1
        result[unique] = _json_value(value)
    return result


def _tabular_record(columns: dict, location: dict, vehicle: str, market: str,
                    inherited_vehicles: list[str], inherited_market: str | None) -> dict | None:
    testimony = _lookup(columns, _FEEDBACK_KEYS, text_only=True)
    feedback = testimony or _lookup(columns, _TEXT_KEYS)
    content = _text(feedback)
    if not content:
        content = "; ".join(f"{name}: {_text(value)}" for name, value in columns.items() if value is not None and _text(value))
    if not content:
        return None
    explicit_vehicle = _lookup(columns, _VEHICLE_KEYS)
    vehicles = _vehicles(explicit_vehicle) or ([_text(explicit_vehicle).upper()] if explicit_vehicle else inherited_vehicles)
    source_market = _lookup(columns, _MARKET_KEYS)
    record_type = "customer_feedback" if testimony else "context"
    if not testimony and any(re.fullmatch(r"volume\d{4}", _key(name)) for name in columns):
        record_type = "sales_aggregate"
    metadata = {"columns": columns, "record_type": record_type}
    for name, keys in {
        "feedback_id": ("id", "feedbackid"),
        "provenance": ("source", "provenance", "sourceurl"),
        "topic": ("vfclevel2name", "topic", "category"),
        "subtopic": ("vfclevel3name", "subtopic"),
        "feedback_type": ("feedbacktype",),
        "sentiment": ("sentiment", "feedbacktype"),
        "engine_type": ("enginetype",),
    }.items():
        value = _lookup(columns, keys)
        if value is not None:
            metadata[name] = value
    raw_date = _lookup(columns, _DATE_KEYS)
    observed_at = _date_value(raw_date)
    if raw_date is not None:
        metadata["date_original"] = raw_date
    metadata["date_status"] = "available" if observed_at else "unparsed" if raw_date is not None else "absent"
    return _base_record(content, location, metadata, vehicle, market, vehicles,
                        _text(source_market) if source_market else inherited_market, observed_at)


def _study_records(rows: list[tuple], sheet: str, header_index: int, vehicle: str, market: str):
    """Read vehicle-column survey matrices, preserving merged country headings."""
    from openpyxl.utils import get_column_letter

    headers = rows[header_index]
    vehicle_columns = {i: _vehicles(value) for i, value in enumerate(headers) if _vehicles(value)}
    countries: dict[int, str] = {}
    for row in rows[:header_index]:
        current = None
        for i in range(max(vehicle_columns) + 1):
            value = row[i] if i < len(row) else None
            label = _text(value).upper()
            if label in _MARKETS or label in {"US", "EU", "CN", "GB", "DE", "ROW"}:
                current = normalize_market(label)
            if current and i in vehicle_columns:
                countries[i] = current
    topic = _text(headers[0])
    topic_row = header_index + 1
    year = re.search(r"(?<!\d)(20\d{2})(?!\d)", sheet)
    for index, row in enumerate(rows[header_index + 1:], header_index + 2):
        if not any(value is not None for value in row):
            continue
        values = [(col, row[col]) for col in vehicle_columns if col < len(row) and row[col] is not None]
        if not values:
            if row[0] is not None:
                topic, topic_row = _text(row[0]), index
            continue
        response = _text(row[0])
        units = _text(row[1]) if 1 not in vehicle_columns and len(row) > 1 else ""
        for col, value in values:
            if not isinstance(value, (int, float)) or isinstance(value, bool):
                continue
            metadata = {"record_type": "study_aggregate", "topic": topic, "response": response,
                        "value": _json_value(value), "units": units or None,
                        "vehicle_label": _text(headers[col]), "columns": _columns(list(headers), row),
                        "date_status": "period_only" if year else "absent"}
            if year:
                metadata["reporting_year"] = int(year.group(1))
            location = {"sheet": sheet, "row": index, "column": col + 1,
                        "cell": f"{get_column_letter(col + 1)}{index}", "header_row": header_index + 1,
                        "topic_row": topic_row}
            content = f"Survey aggregate — {topic}; {response}: {value}"
            if units:
                content += f" ({units})"
            yield _base_record(content, location, metadata, vehicle, market,
                               vehicle_columns[col], countries.get(col) or _sheet_market(sheet))


def _xlsx(data: bytes, filename: str, vehicle: str, market: str) -> list[dict]:
    from openpyxl import load_workbook

    try:
        with ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_MEMBERS or sum(item.file_size for item in members) > MAX_XLSX_UNCOMPRESSED:
                raise ValueError("Workbook exceeds decompression limits")
            if any(item.flag_bits & 1 for item in members):
                raise ValueError("Encrypted workbook archives are unsupported")
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except (BadZipFile, OSError, KeyError) as exc:
        raise ValueError("Unreadable XLSX workbook") from exc
    result = _Records()
    total_cells = total_rows = 0
    try:
        for sheet in workbook:
            # Do not trust worksheet dimensions: they may underreport actual rows.
            if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                raise ValueError("Workbook dimensions exceed import limits")
            if (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_CELLS:
                raise ValueError("Workbook dimensions exceed import limits")
            sheet.reset_dimensions()
            rows = []
            for row in sheet.iter_rows(values_only=True):
                total_rows += 1
                total_cells += len(row)
                if total_rows > MAX_ROWS or total_cells > MAX_CELLS or len(row) > MAX_COLUMNS:
                    raise ValueError("Workbook exceeds row or cell limits")
                rows.append(row)
            nonempty = [(i, row) for i, row in enumerate(rows) if any(value is not None for value in row)]
            if not nonempty:
                continue
            first_header_keys = {_key(value) for value in nonempty[0][1] if value is not None}
            is_tabular = bool(first_header_keys.intersection(_TEXT_KEYS + _VEHICLE_KEYS + _MARKET_KEYS))
            study_header = None if is_tabular else next((
                i for i, row in nonempty[:20]
                if any(_vehicles(value) and (
                    re.fullmatch(r"[FG]\d{2}", _text(value), re.I)
                    or re.match(r"BMW\s+.*Series", _text(value), re.I)
                ) for value in row[1:])
            ), None)
            if study_header is not None:
                result.extend(_study_records(rows, sheet.title, study_header, vehicle, market))
                continue
            header_index, header = nonempty[0]
            inherited = _vehicles(sheet.title) or _vehicles(filename)
            for index, row in enumerate(rows[header_index + 1:], header_index + 2):
                if not any(value is not None for value in row):
                    continue
                record = _tabular_record(_columns(list(header), row), {"sheet": sheet.title, "row": index, "header_row": header_index + 1}, vehicle, market, inherited, _sheet_market(sheet.title))
                if record:
                    result.append(record)
    finally:
        workbook.close()
    return result


def _decode(data: bytes) -> str:
    try:
        encoding = "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        text = data.decode(encoding)
    except UnicodeError as exc:
        raise ValueError("Text input must use UTF-8 or BOM-marked UTF-16") from exc
    if "\x00" in text:
        raise ValueError("Text input contains binary data")
    return text


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.output: list[str] = []
        self.hidden: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript", "template"}:
            self.hidden.append(tag)
        if not self.hidden and tag in {"td", "th"}:
            self.output.append(" ")
        if not self.hidden and tag in {"p", "div", "br", "li", "h1", "h2", "h3", "tr", "article", "section"}:
            self.output.append("\n")

    def handle_endtag(self, tag):
        if self.hidden and tag == self.hidden[-1]:
            self.hidden.pop()
        elif not self.hidden and tag in {"p", "div", "li", "h1", "h2", "h3", "tr", "article", "section"}:
            self.output.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            self.output.append(data)


def parse_html(text: str) -> str:
    """Extract readable text; never fetch URLs or execute source markup."""
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError("HTML exceeds text limits")
    parser = _HTMLText()
    parser.feed(text)
    parser.close()
    return "\n".join(line for line in (re.sub(r"\s+", " ", line).strip() for line in "".join(parser.output).splitlines()) if line)


def _pdf(data: bytes, filename: str, vehicle: str, market: str) -> list[dict]:
    from pypdf import PdfReader, apply_configuration

    # pypdf's configuration uses contextvars, so concurrent imports keep their
    # own budgets. Bound decompression before extract_text allocates its data.
    with apply_configuration(
        maximum_declared_stream_length=MAX_PDF_STREAM_BYTES,
        array_based_stream_maximum_output_length=MAX_PDF_STREAM_BYTES,
        zlib_maximum_output_length=MAX_PDF_STREAM_BYTES,
        lzw_maximum_output_length=MAX_PDF_STREAM_BYTES,
        run_length_maximum_output_length=MAX_PDF_STREAM_BYTES,
        brotli_maximum_output_length=MAX_PDF_STREAM_BYTES,
        stream_decoding_work_maximum_length=MAX_PDF_STREAM_BYTES * 4,
        page_tree_maximum_entries=MAX_PDF_PAGES,
    ):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("PDF requires a password")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError("PDF exceeds page limit")
        result = _Records()
        for index, page in enumerate(reader.pages, 1):
            content = (page.extract_text() or "").strip()
            if content:
                result.append(_base_record(content, {"page": index}, {"record_type": "document_context", "date_status": "absent"}, vehicle, market, _vehicles(filename)))
    return result


def _csv(text: str, filename: str, vehicle: str, market: str) -> list[dict]:
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect=dialect, strict=True)
    header = next(reader, None)
    if not header:
        return []
    if len(header) > MAX_COLUMNS:
        raise ValueError("CSV exceeds column limit")
    result = _Records()
    total_cells = len(header)
    line_end = reader.line_num
    for number, row in enumerate(reader, 2):
        line_start, line_end = line_end + 1, reader.line_num
        total_cells += len(row)
        if number > MAX_ROWS or len(row) > MAX_COLUMNS or total_cells > MAX_CELLS:
            raise ValueError("CSV exceeds row, column or cell limits")
        if not any(value.strip() for value in row):
            continue
        if len(row) != len(header):
            raise ValueError(f"CSV record {number} has a different number of fields than the header")
        record = _tabular_record(_columns(header, row), {"row": number, "line_start": line_start, "line_end": line_end, "header_row": 1}, vehicle, market, _vehicles(filename), None)
        if record:
            result.append(record)
    return result


def _json(text: str, filename: str, vehicle: str, market: str) -> list[dict]:
    def unique_object(pairs):
        result = {}
        for key, value in pairs:
            if key in result:
                raise ValueError(f"Duplicate JSON key: {key}")
            result[key] = value
        return result

    parsed = json.loads(text, object_pairs_hook=unique_object,
                        parse_constant=lambda value: (_ for _ in ()).throw(ValueError(f"Invalid JSON number: {value}")))
    pointer_prefix = "/"
    if isinstance(parsed, dict) and "records" in parsed:
        parsed = parsed["records"]
        pointer_prefix = "/records/"
    elif isinstance(parsed, dict):
        parsed = [parsed]
        pointer_prefix = None
    if not isinstance(parsed, list) or len(parsed) > MAX_RECORDS:
        raise ValueError("JSON must contain an object or a bounded list of records")
    result = _Records()
    for index, item in enumerate(parsed):
        if not isinstance(item, dict):
            raise ValueError("JSON records must be objects")
        if len(item) > MAX_COLUMNS:
            raise ValueError("JSON record exceeds field limit")
        location = {"record": index + 1, "json_pointer": f"{pointer_prefix}{index}" if pointer_prefix else ""}
        record = _tabular_record(_json_value(item), location, vehicle, market, _vehicles(filename), None)
        if record:
            result.append(record)
    return result


def parse_file(filename: str, data: bytes, vehicle: str = "G60", market: str = "US") -> list[dict]:
    """Return records in the requested scope, or [] when all records are elsewhere.

    Raises ValueError for empty, unsupported, malformed or excessive input.
    Caller must preserve original bytes and add file/source identity to locations.
    """
    if not isinstance(data, bytes) or not data or len(data) > MAX_FILE_BYTES:
        raise ValueError("File must contain between 1 byte and 32 MiB")
    if not all(isinstance(value, str) and value.strip() for value in (filename, vehicle, market)):
        raise ValueError("Filename, vehicle and market are required")
    vehicle, market = vehicle.strip().upper(), normalize_market(market)
    filename = PurePath(filename.replace("\\", "/")).name
    suffix = PurePath(filename).suffix.lower()
    try:
        if suffix == ".xlsx":
            result = _xlsx(data, filename, vehicle, market)
        elif suffix == ".pdf":
            result = _pdf(data, filename, vehicle, market)
        elif suffix == ".csv":
            result = _csv(_decode(data), filename, vehicle, market)
        elif suffix == ".json":
            result = _json(_decode(data), filename, vehicle, market)
        elif suffix in {".txt", ".md", ".html", ".htm"}:
            text = _decode(data)
            content = parse_html(text) if suffix in {".html", ".htm"} else text.strip()
            location = {"line_start": 1, "line_end": len(text.splitlines())}
            result = [_base_record(content, location, {"record_type": "document_context", "date_status": "absent"}, vehicle, market, _vehicles(filename))] if content else []
        else:
            raise ValueError(f"Unsupported file type: {suffix or '(none)'}")
        if not result:
            raise ValueError("File contains no readable records")
        if len(result) > MAX_RECORDS or sum(len(item["content"]) for item in result) > MAX_TEXT_CHARS:
            raise ValueError("File exceeds record or extracted text limits")
        return [item for item in result if item["vehicle"] == vehicle and item["market"] == market]
    except ValueError:
        raise
    except Exception as exc:
        # Third-party parsers expose many malformed-file exception classes.
        raise ValueError(f"Unable to read {suffix or 'file'} input: {type(exc).__name__}") from exc
