# /// script
# dependencies = ["openpyxl", "pymupdf", "python-pptx"]
# ///
"""Create searchable derivatives while retaining untouched source files in archive/."""
from pathlib import Path
import json
import openpyxl
import pymupdf
from pptx import Presentation

root=Path(__file__).resolve().parents[1]
out=root/'archive/extracted'; out.mkdir(exist_ok=True,parents=True)
inputs=sorted((root/'archive/drive').rglob('*'))+sorted((root/'archive/documents').glob('*'))
for p in inputs:
    if p.suffix.lower()=='.pdf':
        with pymupdf.open(p) as doc:
            text='\n\n'.join(f'--- Page {i+1} ---\n'+page.get_text() for i,page in enumerate(doc))
            (out/(p.stem+'.txt')).write_text(text)
            links=[{'page':i+1,'links':page.get_links()} for i,page in enumerate(doc) if page.get_links()]
            (out/(p.stem+'.links.json')).write_text(json.dumps(links,default=str,indent=2))
            for i,page in enumerate(doc):
                if len(page.get_text().strip())<20:
                    page.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5)).save(out/f'{p.stem}-page-{i+1}.png')
            print(p.name, len(doc),'pages')
    elif p.suffix.lower()=='.pptx':
        prs=Presentation(p); chunks=[]
        def texts(shapes):
            for sh in shapes:
                if sh.has_text_frame: yield sh.text
                if sh.has_table:
                    for row in sh.table.rows: yield '\t'.join(c.text for c in row.cells)
                if hasattr(sh,'shapes'): yield from texts(sh.shapes)
        for i,s in enumerate(prs.slides):
            chunks.extend([f'--- Slide {i+1} ---',*texts(s.shapes)])
            if s.has_notes_slide: chunks.append('NOTES: '+s.notes_slide.notes_text_frame.text)
        (out/(p.stem+'.txt')).write_text('\n\n'.join(chunks))
        print(p.name,len(prs.slides),'slides')
    elif p.suffix.lower()=='.xlsx':
        wb=openpyxl.load_workbook(p,data_only=False)
        cached=openpyxl.load_workbook(p,data_only=True)
        data={s.title:list(s.values) for s in wb}
        (out/(p.stem+'.json')).write_text(json.dumps(data,ensure_ascii=False,default=str,indent=2))
        values={s.title:list(s.values) for s in cached}
        (out/(p.stem+'.cached-values.json')).write_text(json.dumps(values,ensure_ascii=False,default=str,indent=2))
        metadata={s.title:{'rows':s.max_row,'columns':s.max_column,'state':s.sheet_state,
            'merged_cells':[str(r) for r in s.merged_cells.ranges],
            'formulas':[{'cell':c.coordinate,'formula':c.value,'cached':cached[s.title][c.coordinate].value}
                for row in s for c in row if c.data_type=='f']} for s in wb}
        (out/(p.stem+'.metadata.json')).write_text(json.dumps(metadata,ensure_ascii=False,default=str,indent=2))
        print(p.name,[(s.title,s.max_row,s.max_column) for s in wb])
